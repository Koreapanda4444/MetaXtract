from __future__ import annotations

import hashlib
import io
import zipfile
from dataclasses import dataclass
from pathlib import Path

from metaxtract.case.paths import resolve_source
from metaxtract.config import MEBIBYTE


MAX_PREVIEW_FILE_BYTES = 64 * MEBIBYTE
MAX_PREVIEW_ARCHIVE_BYTES = 64 * MEBIBYTE
MAX_PREVIEW_ARCHIVE_MEMBERS = 2_000
MAX_PREVIEW_COMPRESSION_RATIO = 200.0
MAX_PREVIEW_IMAGE_PIXELS = 40_000_000
MAX_PREVIEW_TEXT_BYTES = 2 * MEBIBYTE
MAX_PREVIEW_TEXT_CHARS = 30_000
MAX_PREVIEW_PDF_PAGES = 3
MAX_PREVIEW_IMAGE_SIZE = (960, 720)

TEXT_SUFFIXES = {
    ".csv",
    ".ini",
    ".json",
    ".log",
    ".md",
    ".py",
    ".toml",
    ".tsv",
    ".txt",
    ".xml",
    ".yaml",
    ".yml",
}
IMAGE_SUFFIXES = {".bmp", ".gif", ".jpeg", ".jpg", ".png", ".tif", ".tiff", ".webp"}


@dataclass(frozen=True)
class PreviewContent:
    kind: str
    title: str
    source_path: str
    source_size: int
    integrity: str
    text: str = ""
    image_png: bytes | None = None
    truncated: bool = False


def resolve_preview_source(
    base: str | Path,
    relative_path: str,
    *,
    expected_file: str | Path | None = None,
) -> Path:
    base_path = Path(base).expanduser()
    if base_path.is_symlink():
        raise ValueError("preview base cannot be a symbolic link")
    resolved_base = base_path.resolve(strict=True)
    if not resolved_base.is_dir():
        raise NotADirectoryError(f"preview base is not a directory: {base_path}")
    source = resolve_source(resolved_base, relative_path)
    if expected_file is not None:
        expected_path = Path(expected_file).expanduser()
        if expected_path.is_symlink():
            raise ValueError("preview source cannot be a symbolic link")
        if source != expected_path.resolve(strict=True):
            raise ValueError("record does not match the scanned source file")
    return source


def _integrity_status(
    path: Path,
    size: int,
    expected_size: int | None,
    expected_sha256: str | None,
) -> str:
    if expected_size is None or not expected_sha256 or len(expected_sha256) != 64:
        return "Original integrity was not available for this record."
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        while chunk := stream.read(MEBIBYTE):
            digest.update(chunk)
    if size == expected_size and digest.hexdigest().casefold() == expected_sha256.casefold():
        return "Original matches the recorded size and SHA-256."
    return "Warning: original no longer matches the recorded size or SHA-256."


def _bounded_text(text: str) -> tuple[str, bool]:
    if len(text) <= MAX_PREVIEW_TEXT_CHARS:
        return text, False
    return f"{text[:MAX_PREVIEW_TEXT_CHARS]}\n\n[preview truncated]", True


def _image_preview(path: Path) -> tuple[bytes, str]:
    from PIL import Image, ImageOps

    with Image.open(path) as image:
        width, height = image.size
        if width < 1 or height < 1 or width * height > MAX_PREVIEW_IMAGE_PIXELS:
            raise ValueError(f"image preview exceeds {MAX_PREVIEW_IMAGE_PIXELS:,} pixels")
        image.seek(0)
        rendered = ImageOps.exif_transpose(image)
        rendered.thumbnail(MAX_PREVIEW_IMAGE_SIZE, Image.Resampling.LANCZOS)
        if rendered.mode not in {"RGB", "RGBA"}:
            rendered = rendered.convert("RGBA" if "transparency" in rendered.info else "RGB")
        output = io.BytesIO()
        rendered.save(output, format="PNG", optimize=True)
        description = (
            f"{image.format or 'Image'} · {width} × {height} pixels · "
            f"preview {rendered.width} × {rendered.height}"
        )
    return output.getvalue(), description


def _pdf_preview(path: Path) -> tuple[str, bool]:
    from pypdf import PdfReader

    reader = PdfReader(str(path), strict=False)
    if reader.is_encrypted:
        try:
            unlocked = bool(reader.decrypt(""))
        except Exception:
            unlocked = False
        if not unlocked:
            return "Encrypted PDF. A password is required to preview its text.", False

    page_count = len(reader.pages)
    sections = [f"PDF pages: {page_count}"]
    for index in range(min(page_count, MAX_PREVIEW_PDF_PAGES)):
        try:
            text = reader.pages[index].extract_text() or ""
        except Exception as exc:
            text = f"[page text unavailable: {type(exc).__name__}]"
        sections.append(f"\n--- Page {index + 1} ---\n{text.strip()}")
        if sum(map(len, sections)) > MAX_PREVIEW_TEXT_CHARS:
            break
    preview, truncated = _bounded_text("\n".join(sections))
    if page_count > MAX_PREVIEW_PDF_PAGES:
        preview += f"\n\n[showing the first {MAX_PREVIEW_PDF_PAGES} pages]"
        truncated = True
    return preview, truncated


def _validate_docx_archive(path: Path) -> None:
    with zipfile.ZipFile(path, "r") as archive:
        members = archive.infolist()
        if len(members) > MAX_PREVIEW_ARCHIVE_MEMBERS:
            raise ValueError("document contains too many archive members")
        total_size = 0
        for member in members:
            total_size += member.file_size
            if total_size > MAX_PREVIEW_ARCHIVE_BYTES:
                raise ValueError("document expands beyond the preview limit")
            compressed = max(member.compress_size, 1)
            if member.file_size / compressed > MAX_PREVIEW_COMPRESSION_RATIO:
                raise ValueError("document contains an unsafe compression ratio")


def _docx_preview(path: Path) -> tuple[str, bool]:
    import docx

    _validate_docx_archive(path)
    document = docx.Document(str(path))
    parts = []
    length = 0

    def append(value: str) -> bool:
        nonlocal length
        value = value.strip()
        if not value:
            return True
        remaining = MAX_PREVIEW_TEXT_CHARS - length
        if remaining <= 0:
            return False
        parts.append(value[:remaining])
        length += min(len(value), remaining)
        return len(value) <= remaining

    complete = True
    for paragraph in document.paragraphs:
        if not append(paragraph.text):
            complete = False
            break
    if complete:
        for table in document.tables:
            for row in table.rows:
                if not append("\t".join(cell.text for cell in row.cells)):
                    complete = False
                    break
            if not complete:
                break
    text = "\n".join(parts)
    if not complete:
        text += "\n\n[preview truncated]"
    return text or "Document contains no previewable text.", not complete


def _plain_text_preview(path: Path) -> tuple[str, bool]:
    with path.open("rb") as stream:
        data = stream.read(MAX_PREVIEW_TEXT_BYTES + 1)
    byte_truncated = len(data) > MAX_PREVIEW_TEXT_BYTES
    text = data[:MAX_PREVIEW_TEXT_BYTES].decode("utf-8-sig", errors="replace")
    text, character_truncated = _bounded_text(text)
    if byte_truncated and not character_truncated:
        text += "\n\n[preview truncated]"
    return text, byte_truncated or character_truncated


def build_record_preview(
    base: str | Path,
    relative_path: str,
    *,
    mime: str = "",
    expected_file: str | Path | None = None,
    expected_size: int | None = None,
    expected_sha256: str | None = None,
) -> PreviewContent:
    source = resolve_preview_source(base, relative_path, expected_file=expected_file)
    source_size = source.stat().st_size
    if source_size > MAX_PREVIEW_FILE_BYTES:
        raise ValueError(f"preview file exceeds {MAX_PREVIEW_FILE_BYTES} bytes")
    integrity = _integrity_status(source, source_size, expected_size, expected_sha256)
    suffix = source.suffix.casefold()
    normalized_mime = mime.casefold()

    if normalized_mime.startswith("image/") or suffix in IMAGE_SUFFIXES:
        try:
            image_png, description = _image_preview(source)
        except (OSError, ValueError):
            raise
        except Exception as exc:
            raise ValueError(f"image preview failed: {type(exc).__name__}") from exc
        return PreviewContent(
            kind="image",
            title=relative_path,
            source_path=str(source),
            source_size=source_size,
            integrity=integrity,
            text=description,
            image_png=image_png,
        )
    if normalized_mime == "application/pdf" or suffix == ".pdf":
        try:
            text, truncated = _pdf_preview(source)
        except (OSError, ValueError):
            raise
        except Exception as exc:
            raise ValueError(f"PDF preview failed: {type(exc).__name__}") from exc
        return PreviewContent(
            kind="text",
            title=relative_path,
            source_path=str(source),
            source_size=source_size,
            integrity=integrity,
            text=text,
            truncated=truncated,
        )
    if suffix == ".docx":
        try:
            text, truncated = _docx_preview(source)
        except (OSError, ValueError, zipfile.BadZipFile):
            raise
        except Exception as exc:
            raise ValueError(f"DOCX preview failed: {type(exc).__name__}") from exc
        return PreviewContent(
            kind="text",
            title=relative_path,
            source_path=str(source),
            source_size=source_size,
            integrity=integrity,
            text=text,
            truncated=truncated,
        )
    if normalized_mime.startswith("text/") or suffix in TEXT_SUFFIXES:
        text, truncated = _plain_text_preview(source)
        return PreviewContent(
            kind="text",
            title=relative_path,
            source_path=str(source),
            source_size=source_size,
            integrity=integrity,
            text=text,
            truncated=truncated,
        )
    raise ValueError("preview is available for images, PDF, DOCX, and text files")
