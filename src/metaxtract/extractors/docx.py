from __future__ import annotations

from typing import Any, Dict, List, Tuple

import docx

from .normalize import as_int, json_safe, normalize_metadata
from ..core.files import PathLike


def extract_docx(path: PathLike) -> Tuple[Dict[str, Any], List[str]]:
    warnings: List[str] = []
    document = docx.Document(str(path))
    props = document.core_properties

    metadata: Dict[str, Any] = {
        "docx_paragraphs": len(document.paragraphs),
        "docx_tables": len(document.tables),
        "docx_sections": len(document.sections),
        "docx_title": props.title or "",
        "docx_author": props.author or "",
        "docx_last_modified_by": props.last_modified_by or "",
        "docx_subject": props.subject or "",
        "docx_keywords": props.keywords or "",
        "docx_category": props.category or "",
        "docx_comments": props.comments or "",
        "docx_language": props.language or "",
        "docx_identifier": props.identifier or "",
        "docx_version": props.version or "",
        "docx_content_status": props.content_status or "",
    }

    for property_name, output_name in (
        ("created", "docx_created"),
        ("modified", "docx_modified"),
        ("last_printed", "docx_last_printed"),
    ):
        value = getattr(props, property_name, None)
        if value is not None:
            metadata[output_name] = json_safe(value)

    revision = as_int(props.revision)
    if revision is not None:
        metadata["docx_revision"] = revision

    return normalize_metadata(metadata), warnings
